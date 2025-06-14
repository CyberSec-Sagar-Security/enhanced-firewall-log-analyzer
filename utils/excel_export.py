"""
Excel Export Utility for SOC Firewall Log Analyzer
Exports alert data to Excel format for SOC analysts
"""

import pandas as pd
from datetime import datetime
import os
from typing import List, Dict, Any
from loguru import logger


class ExcelExporter:
    """Handles Excel export of alert data"""
    
    def __init__(self, output_dir: str = "logs"):
        """
        Initialize Excel exporter
        
        Args:
            output_dir: Directory to save Excel files
        """
        self.output_dir = output_dir
        
        # Ensure output directory exists
        os.makedirs(output_dir, exist_ok=True)
    
    def export_alerts(self, alerts: List[Dict[str, Any]], filename: str = None) -> str:
        """
        Export alerts to Excel file
        
        Args:
            alerts: List of alert dictionaries
            filename: Optional custom filename
            
        Returns:
            Path to created Excel file
        """
        if not alerts:
            logger.warning("No alerts to export to Excel")
            return None
        
        # Generate filename if not provided
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"firewall_alerts_{timestamp}.xlsx"
        
        # Ensure .xlsx extension
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        
        filepath = os.path.join(self.output_dir, filename)
        
        try:
            # Convert alerts to DataFrame
            df = self._prepare_dataframe(alerts)
            
            # Create Excel writer with multiple sheets
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                # Main alerts sheet
                df.to_excel(writer, sheet_name='All Alerts', index=False)
                
                # Summary by threat type
                summary_df = self._create_summary_sheet(df)
                summary_df.to_excel(writer, sheet_name='Summary', index=False)
                
                # High priority alerts only
                high_priority_df = df[df['Severity'].isin(['Critical', 'High'])]
                if not high_priority_df.empty:
                    high_priority_df.to_excel(writer, sheet_name='High Priority', index=False)
                
                # Format the sheets
                self._format_excel_sheets(writer, df)
            
            logger.info(f"✅ Excel export completed: {filepath}")
            return filepath
            
        except Exception as e:
            logger.error(f"❌ Excel export failed: {e}")
            return None
    
    def _prepare_dataframe(self, alerts: List[Dict[str, Any]]) -> pd.DataFrame:
        """
        Convert alerts to pandas DataFrame with proper columns
        
        Args:
            alerts: List of alert dictionaries
            
        Returns:
            Formatted DataFrame
        """
        # Extract relevant fields from alerts
        rows = []
        for alert in alerts:
            row = {
                'Timestamp': alert.get('timestamp', ''),
                'Threat Type': alert.get('threat_type', ''),
                'Severity': alert.get('severity', ''),
                'Source IP': alert.get('source_ip', ''),
                'Destination IP': alert.get('destination_ip', ''),
                'Port': alert.get('port', ''),
                'Protocol': alert.get('protocol', ''),
                'Description': alert.get('description', ''),
                'Details': alert.get('details', ''),
                'Action': alert.get('action', ''),
                'Interface': alert.get('interface', ''),
                'Count': alert.get('count', 1)
            }
            rows.append(row)
        
        df = pd.DataFrame(rows)
        
        # Convert timestamp to datetime for better sorting
        if 'Timestamp' in df.columns:
            df['Timestamp'] = pd.to_datetime(df['Timestamp'], errors='coerce')
        
        # Sort by timestamp (newest first)
        df = df.sort_values('Timestamp', ascending=False)
        
        return df
    
    def _create_summary_sheet(self, df: pd.DataFrame) -> pd.DataFrame:
        """
        Create summary statistics sheet
        
        Args:
            df: Main alerts DataFrame
            
        Returns:
            Summary DataFrame
        """
        summary_data = []
        
        # Total alerts
        summary_data.append({
            'Metric': 'Total Alerts',
            'Count': len(df),
            'Details': 'All detected threats'
        })
        
        # By severity
        if 'Severity' in df.columns:
            severity_counts = df['Severity'].value_counts()
            for severity, count in severity_counts.items():
                summary_data.append({
                    'Metric': f'{severity} Severity',
                    'Count': count,
                    'Details': f'Alerts with {severity.lower()} severity'
                })
        
        # By threat type
        if 'Threat Type' in df.columns:
            threat_counts = df['Threat Type'].value_counts()
            for threat_type, count in threat_counts.items():
                summary_data.append({
                    'Metric': threat_type,
                    'Count': count,
                    'Details': f'Instances of {threat_type.lower()}'
                })
        
        # Top source IPs
        if 'Source IP' in df.columns:
            top_sources = df['Source IP'].value_counts().head(5)
            for ip, count in top_sources.items():
                summary_data.append({
                    'Metric': f'Top Source: {ip}',
                    'Count': count,
                    'Details': 'Alerts from this source IP'
                })
        
        return pd.DataFrame(summary_data)
    
    def _format_excel_sheets(self, writer, df: pd.DataFrame):
        """
        Apply formatting to Excel sheets
        
        Args:
            writer: Excel writer object
            df: Main DataFrame for reference
        """
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            
            # Format All Alerts sheet
            workbook = writer.book
            if 'All Alerts' in workbook.sheetnames:
                worksheet = workbook['All Alerts']
                
                # Header formatting
                header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                header_font = Font(color="FFFFFF", bold=True)
                
                for col in range(1, len(df.columns) + 1):
                    cell = worksheet.cell(row=1, column=col)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")
                
                # Auto-adjust column widths
                for col in range(1, len(df.columns) + 1):
                    column_letter = get_column_letter(col)
                    max_length = 0
                    for row in range(1, min(worksheet.max_row + 1, 100)):  # Check first 100 rows
                        try:
                            cell_value = str(worksheet[f"{column_letter}{row}"].value)
                            max_length = max(max_length, len(cell_value))
                        except:
                            pass
                    
                    # Set column width (max 50 characters)
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Color-code severity cells
                if 'Severity' in df.columns:
                    severity_col = df.columns.get_loc('Severity') + 1
                    severity_colors = {
                        'Critical': PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid"),
                        'High': PatternFill(start_color="FF6600", end_color="FF6600", fill_type="solid"),
                        'Medium': PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid"),
                        'Low': PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
                    }
                    
                    for row in range(2, worksheet.max_row + 1):
                        cell = worksheet.cell(row=row, column=severity_col)
                        if cell.value in severity_colors:
                            cell.fill = severity_colors[cell.value]
                            if cell.value in ['Critical', 'High']:
                                cell.font = Font(color="FFFFFF", bold=True)
        
        except ImportError:
            logger.warning("openpyxl formatting features not available")
        except Exception as e:
            logger.warning(f"Excel formatting failed: {e}")


def export_alerts_to_excel(alerts: List[Dict[str, Any]], output_dir: str = "logs") -> str:
    """
    Convenience function to export alerts to Excel
    
    Args:
        alerts: List of alert dictionaries
        output_dir: Directory to save Excel file
        
    Returns:
        Path to created Excel file
    """
    exporter = ExcelExporter(output_dir)
    return exporter.export_alerts(alerts)
