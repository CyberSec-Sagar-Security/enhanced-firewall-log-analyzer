#!/usr/bin/env python3
"""
Excel Report Generator
Creates comprehensive Excel reports from firewall log analysis
Author: SOC Team
"""

import os
import json
from datetime import datetime
from typing import Dict, List
import csv


class ExcelReportGenerator:
    """Generates Excel reports from firewall analysis results"""
    
    def __init__(self, output_dir: str = "reports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
        
        # Try to import openpyxl, fall back to CSV if not available
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils.dataframe import dataframe_to_rows
            self.excel_available = True
            self.openpyxl = openpyxl
            self.Font = Font
            self.PatternFill = PatternFill
            self.Alignment = Alignment
        except ImportError:
            self.excel_available = False
            print("⚠️  openpyxl not available - will generate CSV reports instead")
    
    def create_comprehensive_report(self, alerts: List[Dict], session_info: Dict) -> str:
        """
        Create comprehensive Excel/CSV report
        
        Args:
            alerts: List of alert dictionaries
            session_info: Analysis session information
            
        Returns:
            Path to created report file
        """
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        if self.excel_available:
            filename = f"Firewall_Analysis_Report_{timestamp}.xlsx"
            filepath = os.path.join(self.output_dir, filename)
            return self._create_excel_report(alerts, session_info, filepath)
        else:
            filename = f"Firewall_Analysis_Report_{timestamp}.csv"
            filepath = os.path.join(self.output_dir, filename)
            return self._create_csv_report(alerts, session_info, filepath)
    
    def _create_excel_report(self, alerts: List[Dict], session_info: Dict, filepath: str) -> str:
        """Create Excel report with multiple sheets"""
        
        wb = self.openpyxl.Workbook()
        
        # Remove default sheet
        wb.remove(wb.active)
        
        # Create sheets
        self._create_summary_sheet(wb, alerts, session_info)
        self._create_alerts_sheet(wb, alerts)
        self._create_threat_analysis_sheet(wb, alerts)
        
        wb.save(filepath)
        return filepath
    
    def _create_summary_sheet(self, wb, alerts: List[Dict], session_info: Dict):
        """Create executive summary sheet"""
        
        ws = wb.create_sheet("Executive Summary", 0)
        
        # Title
        ws['A1'] = "🔥 SOC FIREWALL LOG ANALYZER - EXECUTIVE REPORT"
        ws['A1'].font = self.Font(size=16, bold=True, color="FF0000")
        
        # Analysis info
        row = 3
        ws[f'A{row}'] = "Analysis Information"
        ws[f'A{row}'].font = self.Font(bold=True, size=12)
        row += 1
        
        analysis_data = [
            ["Log File", session_info.get('log_file', 'N/A')],
            ["Parser Type", session_info.get('parser_type', 'N/A')],
            ["Analysis Time", datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
            ["Total Log Entries", session_info.get('total_entries', 0)],
            ["Analysis Duration", session_info.get('duration', 'N/A')]
        ]
        
        for data in analysis_data:
            ws[f'A{row}'] = data[0]
            ws[f'B{row}'] = data[1]
            ws[f'A{row}'].font = self.Font(bold=True)
            row += 1
        
        # Alert summary
        row += 2
        ws[f'A{row}'] = "Alert Summary"
        ws[f'A{row}'].font = self.Font(bold=True, size=12)
        row += 1
        
        # Count alerts by severity
        severity_counts = {}
        type_counts = {}
        
        for alert in alerts:
            severity = alert.get('severity', 'UNKNOWN')
            alert_type = alert.get('type', 'UNKNOWN')
            severity_counts[severity] = severity_counts.get(severity, 0) + 1
            type_counts[alert_type] = type_counts.get(alert_type, 0) + 1
        
        # Severity breakdown
        ws[f'A{row}'] = "Severity"
        ws[f'B{row}'] = "Count"
        ws[f'A{row}'].font = self.Font(bold=True)
        ws[f'B{row}'].font = self.Font(bold=True)
        row += 1
        
        for severity, count in sorted(severity_counts.items()):
            ws[f'A{row}'] = severity
            ws[f'B{row}'] = count
            
            # Color code by severity
            if severity == 'CRITICAL':
                ws[f'A{row}'].fill = self.PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
                ws[f'A{row}'].font = self.Font(color="FFFFFF", bold=True)
            elif severity == 'HIGH':
                ws[f'A{row}'].fill = self.PatternFill(start_color="FF6600", end_color="FF6600", fill_type="solid")
                ws[f'A{row}'].font = self.Font(color="FFFFFF", bold=True)
            elif severity == 'MEDIUM':
                ws[f'A{row}'].fill = self.PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            
            row += 1
        
        # Top threats
        row += 2
        ws[f'A{row}'] = "Top Threat Types"
        ws[f'A{row}'].font = self.Font(bold=True, size=12)
        row += 1
        
        ws[f'A{row}'] = "Threat Type"
        ws[f'B{row}'] = "Count"
        ws[f'A{row}'].font = self.Font(bold=True)
        ws[f'B{row}'].font = self.Font(bold=True)
        row += 1
        
        for threat_type, count in sorted(type_counts.items(), key=lambda x: x[1], reverse=True)[:10]:
            ws[f'A{row}'] = threat_type.replace('_', ' ').title()
            ws[f'B{row}'] = count
            row += 1
        
        # Recommendations
        row += 2
        ws[f'A{row}'] = "Key Recommendations"
        ws[f'A{row}'].font = self.Font(bold=True, size=12)
        row += 1
        
        recommendations = self._generate_recommendations(alerts)
        for i, rec in enumerate(recommendations[:5], 1):
            ws[f'A{row}'] = f"{i}. {rec}"
            row += 1
        
        # Auto-size columns
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_alerts_sheet(self, wb, alerts: List[Dict]):
        """Create detailed alerts sheet"""
        
        ws = wb.create_sheet("Detailed Alerts")
        
        # Headers
        headers = [
            "Alert ID", "Severity", "Type", "Title", "Description", 
            "Timestamp", "Source IP", "Destination IP", "Port", 
            "Recommendation", "Risk Level"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.Font(bold=True)
            cell.fill = self.PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Alert data
        for row, alert in enumerate(alerts, 2):
            data = [
                alert.get('alert_id', f"ALERT-{row-1}"),
                alert.get('severity', ''),
                alert.get('type', '').replace('_', ' ').title(),
                alert.get('title', ''),
                alert.get('description', '')[:100] + "..." if len(alert.get('description', '')) > 100 else alert.get('description', ''),
                alert.get('timestamp', ''),
                alert.get('source_ip', ''),
                alert.get('destination_ip', ''),
                alert.get('destination_port', ''),
                alert.get('recommendation', '')[:150] + "..." if len(alert.get('recommendation', '')) > 150 else alert.get('recommendation', ''),
                alert.get('risk_level', alert.get('severity', ''))
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=str(value) if value is not None else "")
                
                # Color code severity
                severity = alert.get('severity', '')
                if severity == 'CRITICAL':
                    cell.fill = self.PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
                elif severity == 'HIGH':
                    cell.fill = self.PatternFill(start_color="FFE6CC", end_color="FFE6CC", fill_type="solid")
                elif severity == 'MEDIUM':
                    cell.fill = self.PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
        
        # Auto-size columns
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 80)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_threat_analysis_sheet(self, wb, alerts: List[Dict]):
        """Create threat analysis sheet"""
        
        ws = wb.create_sheet("Threat Analysis")
        
        # Threat statistics
        ws['A1'] = "Threat Analysis & Statistics"
        ws['A1'].font = self.Font(size=14, bold=True)
        
        row = 3
        
        # Top source IPs
        source_ips = {}
        dest_ports = {}
        
        for alert in alerts:
            src_ip = alert.get('source_ip', '')
            dst_port = alert.get('destination_port', '')
            
            if src_ip and src_ip != 'TCP':  # Filter out parsing errors
                source_ips[src_ip] = source_ips.get(src_ip, 0) + 1
            
            if dst_port and str(dst_port).isdigit():
                dest_ports[dst_port] = dest_ports.get(dst_port, 0) + 1
        
        # Top attacking IPs
        ws[f'A{row}'] = "Top Attacking Source IPs"
        ws[f'A{row}'].font = self.Font(bold=True)
        row += 1
        
        ws[f'A{row}'] = "Source IP"
        ws[f'B{row}'] = "Alert Count"
        ws[f'A{row}'].font = self.Font(bold=True)
        ws[f'B{row}'].font = self.Font(bold=True)
        row += 1
        
        for ip, count in sorted(source_ips.items(), key=lambda x: x[1], reverse=True)[:10]:
            ws[f'A{row}'] = ip
            ws[f'B{row}'] = count
            row += 1
        
        # Top targeted ports
        row += 2
        ws[f'A{row}'] = "Most Targeted Ports"
        ws[f'A{row}'].font = self.Font(bold=True)
        row += 1
        
        ws[f'A{row}'] = "Port"
        ws[f'B{row}'] = "Alert Count"
        ws[f'A{row}'].font = self.Font(bold=True)
        ws[f'B{row}'].font = self.Font(bold=True)
        row += 1
        
        for port, count in sorted(dest_ports.items(), key=lambda x: x[1], reverse=True)[:10]:
            ws[f'A{row}'] = port
            ws[f'B{row}'] = count
            row += 1
        
        # Auto-size columns
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_csv_report(self, alerts: List[Dict], session_info: Dict, filepath: str) -> str:
        """Create CSV report as fallback"""
        
        with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
            fieldnames = [
                'Alert_ID', 'Severity', 'Type', 'Title', 'Description',
                'Timestamp', 'Source_IP', 'Destination_IP', 'Port',
                'Recommendation', 'Risk_Level'
            ]
            
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
            writer.writeheader()
            
            for i, alert in enumerate(alerts, 1):
                writer.writerow({
                    'Alert_ID': alert.get('alert_id', f"ALERT-{i}"),
                    'Severity': alert.get('severity', ''),
                    'Type': alert.get('type', '').replace('_', ' ').title(),
                    'Title': alert.get('title', ''),
                    'Description': alert.get('description', ''),
                    'Timestamp': alert.get('timestamp', ''),
                    'Source_IP': alert.get('source_ip', ''),
                    'Destination_IP': alert.get('destination_ip', ''),
                    'Port': alert.get('destination_port', ''),
                    'Recommendation': alert.get('recommendation', ''),
                    'Risk_Level': alert.get('risk_level', alert.get('severity', ''))
                })
        
        return filepath
    
    def _generate_recommendations(self, alerts: List[Dict]) -> List[str]:
        """Generate executive recommendations based on alerts"""
        
        recommendations = []
        
        # Count alert types
        type_counts = {}
        severity_counts = {}
        
        for alert in alerts:
            alert_type = alert.get('type', '')
            severity = alert.get('severity', '')
            type_counts[alert_type] = type_counts.get(alert_type, 0) + 1
            severity_counts[severity] = severity_counts.get(severity, 0) + 1
        
        # Generate specific recommendations
        if severity_counts.get('CRITICAL', 0) > 0:
            recommendations.append("URGENT: Critical threats detected - implement emergency response procedures")
        
        if type_counts.get('REVERSE_SHELL', 0) > 0:
            recommendations.append("Investigate hosts for malware infection - reverse shells detected")
        
        if type_counts.get('BRUTE_FORCE', 0) > 5:
            recommendations.append("Implement stronger authentication controls - multiple brute force attacks")
        
        if type_counts.get('SUSPICIOUS_PORT', 0) > 10:
            recommendations.append("Review firewall rules - extensive suspicious port activity")
        
        if type_counts.get('OFF_HOURS_ACTIVITY', 0) > 0:
            recommendations.append("Monitor off-hours activity - potential insider threats or automated attacks")
        
        # Default recommendations
        if not recommendations:
            recommendations.extend([
                "Continue monitoring for emerging threats",
                "Review and update security policies regularly",
                "Implement network segmentation for critical assets"
            ])
        
        return recommendations


def install_excel_support():
    """Install openpyxl for Excel support"""
    
    try:
        import subprocess
        import sys
        
        print("📦 Installing Excel support (openpyxl)...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
        print("✅ Excel support installed successfully")
        return True
    except Exception as e:
        print(f"❌ Failed to install Excel support: {e}")
        return False
